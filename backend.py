"""
backend.py - Data layer for the Climate Risk (Heatwave / Heat Index) Explorer
==============================================================================
All I/O and data-wrangling lives here, independent of Streamlit or plotting,
so it can be unit-tested and reused (e.g. in a notebook) without pulling in
the UI. Nothing in this module imports streamlit or matplotlib.

Data model
----------
Each scenario workbook (Historical, SSP2-4.5, SSP3-7.0, SSP5-8.5, ...) has:
  - one sheet per State/UT, long format: Latitude, Longitude, Year, <indicators...>
  - a 'Metadata' sheet: a run-metadata block (Field/Value pairs) followed by a
    per-indicator table (Variable Name/Full Name/Definition/Units/Calculation Method)

points_mapped.csv carries the authoritative, scenario-independent District/State
attribution per grid node (from an exact point-in-polygon join against the
official district/state boundary shapefiles; 'match' = 'point' for an exact
join or 'overlap' for a nearest-district fallback on coastal/border cells
where no polygon reached the grid centroid).
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Optional

import numpy as np
import pandas as pd
import openpyxl

EARTH_RADIUS_KM = 6371.0088          # IUGG mean radius, for haversine distance
EXPECTED_GRID_SPACING_DEG = 0.25     # nominal resolution, sanity-checked at load time
SNAP_WARN_KM = EXPECTED_GRID_SPACING_DEG * np.sqrt(2) * 111.0  # ~ one grid-cell diagonal


def _cell(row, i):
    """Safely index into a spreadsheet row that may be a short/empty tuple or
    list. openpyxl's read-only mode can return a genuinely empty `[]` (not a
    None-padded tuple) for some blank rows; plain row[i] indexing crashes on
    those. Every raw-row access in this module goes through this helper."""
    return row[i] if i < len(row) else None


@dataclass
class Scenario:
    """One loaded indicator workbook."""
    label: str                      # e.g. "Projection - SSP370 (2015-2037)"
    source_name: str                # original filename, for diagnostics
    run_metadata: dict
    indicator_metadata: dict        # {var_name: {Full Name, Definition, Units, Calculation Method}}
    data: pd.DataFrame              # Sheet_State, Latitude, Longitude, Year, <indicators...>
    year_min: int
    year_max: int

    @property
    def indicators(self) -> list[str]:
        return list(self.indicator_metadata.keys())


def parse_scenario_workbook(file_bytes: bytes, source_name: str) -> Scenario:
    """Parse one indicator workbook. Raises ValueError with a clear message if
    the file isn't shaped like one of these workbooks (missing 'Metadata'
    sheet, no 'Variable Name' header row, etc.) - callers should catch this
    and skip the file rather than crash the whole app."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"Not a readable .xlsx file: {e}") from e

    if "Metadata" not in wb.sheetnames:
        raise ValueError("No 'Metadata' sheet found.")
    ws = wb["Metadata"]
    rows = list(ws.iter_rows(values_only=True))
    try:
        header_row_idx = next(i for i, r in enumerate(rows) if _cell(r, 0) == "Variable Name")
    except StopIteration:
        raise ValueError("'Metadata' sheet has no 'Variable Name' header row.")
    header = rows[header_row_idx]

    indicator_metadata = {}
    for r in rows[header_row_idx + 1:]:
        name = _cell(r, 0)
        if name is None:
            continue
        indicator_metadata[name] = {
            _cell(header, j + 1): _cell(r, j + 1) for j in range(len(header) - 1)
        }
    if not indicator_metadata:
        raise ValueError("No indicator rows found under the 'Variable Name' header.")

    run_metadata = {
        _cell(r, 0): _cell(r, 1) for r in rows[:header_row_idx] if _cell(r, 0) is not None
    }

    frames = []
    for sheet_name in wb.sheetnames:
        if sheet_name == "Metadata":
            continue
        sheet_ws = wb[sheet_name]
        rows_iter = sheet_ws.iter_rows(values_only=True)
        try:
            sheet_header = next(rows_iter)
        except StopIteration:
            continue  # genuinely empty sheet
        # Some sheets carry trailing blank header cells (leftover formatting/
        # merged-cell artifacts, not real columns) - trim them so the header
        # doesn't end up with duplicate None column names, which breaks
        # pandas' column alignment when concatenating across sheets.
        while sheet_header and sheet_header[-1] is None:
            sheet_header = sheet_header[:-1]
        n_cols = len(sheet_header)
        data_rows = []
        for r in rows_iter:
            if _cell(r, 0) is None:
                continue
            # openpyxl can return a shorter tuple for a row than the header's
            # width when that row's trailing cells were never written (e.g. a
            # blank/NaN indicator value at the end) - pad with None so every
            # row matches the header length before building the DataFrame.
            if len(r) < n_cols:
                r = tuple(r) + (None,) * (n_cols - len(r))
            elif len(r) > n_cols:
                r = r[:n_cols]
            data_rows.append(r)
        if not data_rows:
            continue
        sheet_df = pd.DataFrame(data_rows, columns=sheet_header)
        sheet_df.insert(0, "Sheet_State", sheet_name)
        frames.append(sheet_df)
    wb.close()

    if not frames:
        raise ValueError("No data rows found in any State/UT sheet.")
    data = pd.concat(frames, ignore_index=True)
    if "Year" not in data.columns:
        raise ValueError("State/UT sheets have no 'Year' column.")
    data["Year"] = data["Year"].astype(int)

    year_min, year_max = int(data["Year"].min()), int(data["Year"].max())
    label = f"{run_metadata.get('Scenario', source_name)} ({year_min}-{year_max})"
    return Scenario(
        label=label, source_name=source_name, run_metadata=run_metadata,
        indicator_metadata=indicator_metadata, data=data,
        year_min=year_min, year_max=year_max,
    )


def check_grid_spacing(data: pd.DataFrame, expected_deg: float = EXPECTED_GRID_SPACING_DEG) -> Optional[str]:
    """Returns a warning string if the data's actual grid spacing doesn't
    match the expected resolution, else None."""
    unique_lats = np.sort(data["Latitude"].unique())
    if len(unique_lats) < 2:
        return None
    actual = float(np.median(np.diff(unique_lats)))
    if not np.isclose(actual, expected_deg, atol=1e-6):
        return f"Median latitude spacing is {actual:.4f} deg, expected {expected_deg} deg."
    return None


def load_mapping(file_bytes: bytes) -> pd.DataFrame:
    """Load points_mapped.csv (District/State attribution per grid node) and
    standardise column names."""
    mapped = pd.read_csv(io.BytesIO(file_bytes), dtype={"DIST_LGD": str})
    mapped = mapped.rename(columns={
        "STATE_UT": "State", "DISTRICT": "District",
        "STATE_LGD": "State_ID", "DIST_LGD": "District_ID",
    })
    required = {"Latitude", "Longitude", "State", "District", "State_ID", "District_ID"}
    missing = required - set(mapped.columns)
    if missing:
        raise ValueError(f"Mapping file is missing expected columns: {sorted(missing)}")

    # Generic cleanup: strip stray whitespace/control characters some CSV
    # exports leave on text fields (confirmed present on 6 Assam district
    # names as trailing \r\n in this file) - always safe, no guessing involved.
    for col in ("State", "District"):
        mapped[col] = mapped[col].astype(str).str.strip()

    mapped["District"] = mapped["District"].replace(KNOWN_DISTRICT_NAME_FIXES)

    if mapped.duplicated(subset=["Latitude", "Longitude"]).any():
        mapped = mapped.drop_duplicates(subset=["Latitude", "Longitude"])
    return mapped


# Known source-data text corruption in points_mapped.csv: the Kannada long
# vowels (a-, i-, u-macron) in these district names were replaced with
# placeholder characters ('<', '\', '#') by whatever encoding/export step
# produced the file - confirmed by inspecting the raw district boundary
# shapefile's .dbf bytes directly (see the state/district mapping notebook).
# Two more (Dehradun, Birbhum) show the same pattern with '@'. All 21 are
# well-established, unambiguous Indian district names; corrected here rather
# than left broken. Please verify against your own reference list and flag
# this to whoever maintains the source shapefile.
KNOWN_DISTRICT_NAME_FIXES = {
    "B<galkot": "Bagalkot",
    "B\\dar": "Bidar",
    "Ball<ri": "Ballari",
    "Belag<vi": "Belagavi",
    "Bengal#ru (Rural)": "Bengaluru (Rural)",
    "Bengal#ru (Urban)": "Bengaluru (Urban)",
    "Ch<mar<janagara": "Chamarajanagara",
    "Chikkaball<pura": "Chikkaballapura",
    "Chikkamagal#ru": "Chikkamagaluru",
    "D<vanagere": "Davanagere",
    "Dh<rw<d": "Dharwad",
    "H<ssan": "Hassan",
    "H<veri": "Haveri",
    "Kol<ra": "Kolar",
    "Mys#ru": "Mysuru",
    "R<managara": "Ramanagara",
    "Raich#r": "Raichur",
    "Tumak#ru": "Tumakuru",
    "Y<dgir": "Yadgir",
    "DEHRAD@N": "DEHRADUN",
    "BIRBH@M": "BIRBHUM",
}


def coverage_gap(data: pd.DataFrame, mapped: pd.DataFrame) -> int:
    """Count of grid points in `data` with no matching row in `mapped`."""
    ind_points = data[["Latitude", "Longitude"]].drop_duplicates()
    merged = ind_points.merge(mapped[["Latitude", "Longitude"]], on=["Latitude", "Longitude"],
                               how="left", indicator=True)
    return int((merged["_merge"] == "left_only").sum())


def build_grid_lookup(data: pd.DataFrame) -> pd.DataFrame:
    """One row per unique grid node with the list of workbook sheets it
    appears under (needed only to retrieve indicator rows for that node)."""
    return (
        data.groupby(["Latitude", "Longitude"])["Sheet_State"]
        .unique()
        .reset_index()
        .rename(columns={"Sheet_State": "Sheets"})
    )


def haversine_km(lat1, lon1, lat2, lon2):
    """Vectorised great-circle distance (km). Used instead of flat Euclidean
    degree-distance since 1 deg of longitude shrinks with latitude across
    India's ~6-37N span."""
    lat1r, lon1r, lat2r, lon2r = map(np.radians, (lat1, lon1, lat2, lon2))
    dlat = lat2r - lat1r
    dlon = lon2r - lon1r
    a = np.sin(dlat / 2) ** 2 + np.cos(lat1r) * np.cos(lat2r) * np.sin(dlon / 2) ** 2
    return 2 * EARTH_RADIUS_KM * np.arcsin(np.sqrt(np.clip(a, 0, 1)))


@dataclass
class GridMatch:
    lat: float
    lon: float
    sheets: list
    distance_km: float
    state: str
    district: str
    state_id: object
    district_id: object
    admin_match: str  # 'point' | 'overlap' | 'Unknown'


def nearest_grid_point(lat_in: float, lon_in: float, grid_lookup: pd.DataFrame,
                        admin_lookup: pd.DataFrame) -> GridMatch:
    """Find the nearest actual grid node to (lat_in, lon_in) and its District/
    State attribution."""
    d = haversine_km(lat_in, lon_in, grid_lookup["Latitude"].values, grid_lookup["Longitude"].values)
    i = int(np.argmin(d))
    row = grid_lookup.iloc[i]
    lat, lon = float(row["Latitude"]), float(row["Longitude"])
    try:
        admin = admin_lookup.loc[(lat, lon)]
        state, district, state_id, district_id, admin_match = (
            admin["State"], admin["District"], admin["State_ID"], admin["District_ID"], admin["match"],
        )
    except KeyError:
        state = district = state_id = district_id = "Unknown"
        admin_match = "Unknown"
    return GridMatch(lat, lon, list(row["Sheets"]), float(d[i]), state, district, state_id, district_id, admin_match)


def point_time_series(data: pd.DataFrame, lat: float, lon: float, indicator: str,
                       year_lo: int, year_hi: int) -> pd.DataFrame:
    """Yearly values of `indicator` at exactly (lat, lon), within [year_lo, year_hi].
    Border grid nodes are duplicated identically across two State sheets (verified
    in the notebook's Section 4); de-duplicate to one row per year."""
    subset = data[(data["Latitude"] == lat) & (data["Longitude"] == lon) &
                  (data["Year"].between(year_lo, year_hi))].sort_values("Year")
    return subset.drop_duplicates(subset=["Year"])[["Year", indicator]].reset_index(drop=True)


# --------------------------------------------------------------------------
# District-level spatial aggregation (via points_mapped.csv attribution)
# --------------------------------------------------------------------------

def list_states(mapped: pd.DataFrame) -> list[str]:
    return sorted(mapped["State"].dropna().unique().tolist())


def list_districts(mapped: pd.DataFrame, state: str) -> list[str]:
    return sorted(mapped.loc[mapped["State"] == state, "District"].dropna().unique().tolist())


def district_grid_values(data: pd.DataFrame, mapped: pd.DataFrame, state: str, district: str,
                          indicator: str, year: Optional[int] = None,
                          year_range: Optional[tuple[int, int]] = None) -> pd.DataFrame:
    """Grid-cell-level values of `indicator` for every node attributed to
    (state, district) in points_mapped.csv, either for a single `year` or
    averaged over `year_range` (inclusive). Returns one row per grid node:
    Latitude, Longitude, value, match (point/overlap fallback flag).

    Note on methodology: this attributes each 0.25 deg grid cell to a single
    district using its centroid's point-in-polygon result (with a nearest-
    district fallback for ~8% of coastal/border cells - see points_mapped.csv's
    'match' column). It is NOT an 'all-touched' polygon-raster clip (which
    would need the district boundary shapefile itself, not just this
    per-point attribution table) - a cell straddling two districts is
    counted for only one of them here. Swap in a true rioxarray/geopandas
    clip against the boundary shapefile if that distinction matters for your
    analysis.
    """
    if (year is None) == (year_range is None):
        raise ValueError("Pass exactly one of `year` or `year_range`.")

    nodes = mapped[(mapped["State"] == state) & (mapped["District"] == district)][
        ["Latitude", "Longitude", "match"]
    ]
    if nodes.empty:
        return nodes.assign(value=[])

    merged = data.merge(nodes, on=["Latitude", "Longitude"], how="inner")
    if year is not None:
        merged = merged[merged["Year"] == year]
        agg = merged.groupby(["Latitude", "Longitude", "match"], as_index=False)[indicator].mean()
    else:
        lo, hi = year_range
        merged = merged[merged["Year"].between(lo, hi)]
        agg = merged.groupby(["Latitude", "Longitude", "match"], as_index=False)[indicator].mean()
    return agg.rename(columns={indicator: "value"})


def district_time_series(data: pd.DataFrame, mapped: pd.DataFrame, state: str, district: str,
                          indicator: str, year_lo: int, year_hi: int) -> pd.DataFrame:
    """District-mean yearly time series: spatial mean of `indicator` across
    every grid node attributed to (state, district), for each year."""
    nodes = mapped[(mapped["State"] == state) & (mapped["District"] == district)][["Latitude", "Longitude"]]
    merged = data.merge(nodes, on=["Latitude", "Longitude"], how="inner")
    merged = merged[merged["Year"].between(year_lo, year_hi)]
    return (
        merged.groupby("Year", as_index=False)[indicator]
        .agg(mean="mean", min="min", max="max", n_cells="count")
    )
